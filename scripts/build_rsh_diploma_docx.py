from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "diploma" / "rsh-diploma-final-with-visuals.md"
OUTPUT_DOCX = ROOT / "docs" / "diploma" / "RSH_Дипломная_работа_с_иллюстрациями.docx"


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("left", "top", "right", "bottom"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_placeholder_box(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(15.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_border(
        cell,
        top={"val": "single", "sz": 8, "space": 0, "color": "808080"},
        bottom={"val": "single", "sz": 8, "space": 0, "color": "808080"},
        left={"val": "single", "sz": 8, "space": 0, "color": "808080"},
        right={"val": "single", "sz": 8, "space": 0, "color": "808080"},
    )
    set_cell_shading(cell, "F7F7F7")
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(title)
    r1.bold = True
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(body)
    r2.italic = True
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 14)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def configure_section(section) -> None:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)


def add_heading_paragraph(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    if level == 1:
        run.bold = True


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.add_run(text)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(13)
    run.bold = True


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.75)
    p.add_run(f"• {text}")


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.75)
    p.add_run(text)


def add_image(doc: Document, path: Path, width_inches: float = 6.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("ДИПЛОМНАЯ РАБОТА")
    r.bold = True
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(18)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(18)
    r2 = p2.add_run(
        "Проектирование и разработка веб-приложения интернет-магазина брендовой одежды\n"
        "с 2D-конструктором кастомизации изделий"
    )
    r2.font.name = "Times New Roman"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r2.font.size = Pt(16)
    r2.bold = True

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(12)
    r3 = p3.add_run("Проект: RSH")
    r3.font.name = "Times New Roman"
    r3._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r3.font.size = Pt(15)
    r3.italic = True

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_before = Pt(240)
    r4 = p4.add_run("Москва, 2026")
    r4.font.name = "Times New Roman"
    r4._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r4.font.size = Pt(14)


def build_additional_visual_placeholders(doc: Document) -> None:
    add_page_break(doc)
    add_heading_paragraph(doc, "Приложение В. Дополнительные иллюстрации для окончательной версии диплома", 1)
    add_body_paragraph(
        doc,
        "В данном приложении подготовлены позиции для дальнейшей замены на финальные изображения, "
        "если потребуется привести диплом к оформлению, полностью совпадающему с методическими рекомендациями кафедры."
    )

    placeholders = [
        ("Рисунок 14 — Интерфейс аналога fashion eCommerce-платформы", "Сюда вставляется скриншот первого аналога из раздела 1.3."),
        ("Рисунок 15 — Интерфейс аналога сервиса кастомизации одежды", "Сюда вставляется скриншот второго аналога из раздела 1.3."),
        ("Рисунок 16 — Сравнительный пример страницы товара аналога", "Сюда вставляется скриншот третьего аналога из раздела 1.3."),
        ("Рисунок 17 — Схема или визуализация базы данных проекта RSH", "Сюда вставляется скриншот БД, Prisma Studio или экспортированная ER-схема."),
        ("Рисунок 18 — Производственная инфраструктура и хостинг проекта", "Сюда вставляется скриншот Vercel / Neon / Upstash или сводной схемы production-развертывания."),
    ]
    for caption, text in placeholders:
        add_caption(doc, caption)
        add_placeholder_box(doc, "Место под иллюстрацию", text)
        add_body_paragraph(
            doc,
            "Иллюстрация может быть заменена на итоговый скриншот или диаграмму без изменения основной структуры текста дипломной работы."
        )


def main() -> None:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)

    build_title_page(doc)
    add_page_break(doc)

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    in_mermaid = False
    pending_diagram_placeholder = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if line == "---":
            continue
        if line.startswith("```mermaid"):
            in_mermaid = True
            pending_diagram_placeholder = True
            continue
        if line.startswith("```") and in_mermaid:
            in_mermaid = False
            if pending_diagram_placeholder:
                add_placeholder_box(
                    doc,
                    "Временная диаграмма",
                    "На данном месте размещается схема проекта RSH. При необходимости она может быть заменена на итоговую диаграмму из draw.io."
                )
                pending_diagram_placeholder = False
            continue
        if in_mermaid:
            continue

        if line.startswith("# "):
            title = line[2:].strip()
            if doc.paragraphs and doc.paragraphs[-1].text.strip():
                add_page_break(doc)
            add_heading_paragraph(doc, title, 1)
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            if title.startswith("Рисунок "):
                add_caption(doc, title)
            else:
                add_heading_paragraph(doc, title, 2)
            continue
        if line.startswith("### "):
            add_heading_paragraph(doc, line[4:].strip(), 3)
            continue
        if line.startswith("![") and "](" in line and line.endswith(")"):
            path_part = line.split("](", 1)[1][:-1]
            if path_part.startswith("file:///"):
                image_path = Path(path_part.replace("file:///", ""))
            else:
                image_path = Path(path_part)
            if image_path.exists():
                add_image(doc, image_path)
            else:
                add_placeholder_box(
                    doc,
                    "Изображение не найдено",
                    f"По указанному пути не найден файл: {image_path}"
                )
            continue
        if re.match(r"^\d+\.\s", line):
            add_numbered(doc, line)
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:].strip())
            continue
        add_body_paragraph(doc, line)

    build_additional_visual_placeholders(doc)
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
